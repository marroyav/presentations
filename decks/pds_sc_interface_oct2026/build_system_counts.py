"""Expand the selected variable patterns into an explicit FD-VD draft case."""
from pathlib import Path
from collections import defaultdict
import hashlib
import json
import re
import sys

sys.path.insert(0, '/home/arroyave/work/daphne-sc-intake-20261014/pythondeps')
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from prepare_evidence import workbook_rows

BASE = Path(__file__).resolve().parent
SOURCE = Path('/home/arroyave/Downloads/PDS_Operational_Variables_v1.1_Draft.xlsx')
OUTPUT = Path('/home/arroyave/Downloads/PDS_Operational_Variables_v1.2_System_Draft.xlsx')
INPUTS = [
 ('boards', 42, 'DAPHNE boards', 'Current local SC–PDS ICD §2.1: 1,344 channels / 32. May rack draft used 40 boards; not the same as 40 mini-racks.'),
 ('channels', 32, 'Used channels per board', 'FD-VD design allocation. Set 40 to count the full implemented register capacity instead.'),
 ('afes', 4, 'Used AFE groups per board', 'Working case: 32 used channels / 8 channels per AFE. Set 5 for full implemented AFE capacity.'),
 ('links', 2, 'Data links per board', 'User direction: two data links per board.'),
 ('cal', 8, 'PDS calibration boxes', 'Current local SC–PDS ICD §2.2 planning quantity.'),
 ('cal_channels', 12, 'Channels per calibration box', 'Existing register map/client supports 12; optional 5-channel subset is not this case.'),
 ('cal_voltages', 9, 'Board-voltage inputs per calibration box', 'Existing register map indices 0..8.'),
 ('pof_boxes', 80, 'PoF boxes', 'May rack draft: 40 mini-racks × 2 boxes.'),
 ('top_racks', 40, 'Top-plate mini-racks', 'May rack draft FD2 MiniRack combined!A1.'),
 ('power_racks', 4, 'PDS power racks', 'May rack draft PD LV!B1 and RACKS allocation; differs from old one-crate emulator placement.'),
 ('top_outlets', 8, 'Outlets per top-plate PDU', 'PX3-5147R vendor submittal rev20250625 p.2; candidate in March deck slide 24.'),
 ('power_outlets', 30, 'Outlets per standard-rack PDU', 'PX3-5170V-E2F5 vendor submittal p.2 and DocDB 32655-v8 p.4.'),
 ('temps', 4, 'Temperature probes per rack', 'DocDB 32655-v8 p.21: four Celsius sensors on reference unit. Repetition over all racks is the draft profile.'),
 ('leaks', 1, 'Leak detectors per rack', 'DocDB 32655-v8 required hardware and rack-profile draft.'),
 ('smoke_hubs', 1, 'Smoke interfaces per rack', 'One 4WTAR-B / DX-PD2C5 assembly per rack in the draft profile.'),
 ('contacts', 5, 'Contact-state fields per smoke interface', 'Retained UDT has five directly sourced OnOff 1..5 Status fields; not five smoke detectors.'),
 ('regulators', 4, 'PMBus regulators per board', 'Selected implemented scope 0..3.'),
 ('fans', 2, 'Tachometers per board', 'Selected implemented scope 0..1.'),
 ('current_adc', 10, 'Current ADC inputs per board', 'Selected scope 0..9 at configured mux; no additional sweep positions invented.'),
 ('buffers', None, 'Total Hermes input-buffer instances', 'Firmware reports n_srcs separately from data links; two links do not establish buffer count.'),
 ('hd_afes', None, 'Populated HD-mezzanine AFE instances', 'Count only applicable installed HD-mezzanine groups; do not apply HD-only fields blindly to optical FD-VD boards.'),
 ('power_points', None, 'Defined WIENER power-consumption points', 'Location and reader are not established. DC output power is not mains input consumption.'),
]
values = {key: value for key, value, _, _ in INPUTS}
refs = {key: f"'System Inputs'!$B${i}" for i, (key, *_) in enumerate(INPUTS, 2)}

def expression(keys, operation='product'):
    parts = [refs[k] for k in keys]
    return ('+' if operation == 'sum' else '*').join(parts)

def multiplier(keys, operation='product'):
    if any(values[k] is None for k in keys):
        return None
    if operation == 'sum':
        return sum(values[k] for k in keys)
    result = 1
    for key in keys:
        result *= values[key]
    return result

rows = []
def add(tab, tag, role, expr, quantity, basis, status):
    rows.append(dict(tab=tab, tag=tag, role=role, expression=expr,
                     quantity=quantity, basis=basis, status=status))

book = workbook_rows(SOURCE)
for tab in ['DAQ', 'OPMON', 'SC', 'PoF', 'SSP', 'WIENER', 'PDU']:
    for r in book[tab][1:]:
        if not r.get('B'):
            continue
        tag, role = r['B'], r['C']
        status = r.get('J' if tab in ['DAQ', 'OPMON'] else 'AC', '')
        if tab in ['DAQ', 'OPMON']:
            scope = r['D']
            if (tab == 'DAQ' and 'V2 only' in r.get('H', '')) or (tab == 'OPMON' and 'V2 published' in r.get('G', '')):
                add(tab, tag, role, '0', 0, 'Alternative V2 implementation excluded from V3 case.', status)
                continue
            keys = ['boards']
            if scope.startswith('Channel '): keys.append('channels')
            elif scope.startswith('AFE '): keys.append('afes')
            elif scope == 'Per Hermes link': keys.append('links')
            elif scope == 'Per Hermes input buffer': keys = ['buffers']
            expr, q = expression(keys), multiplier(keys)
            basis = scope + '; selected V3 case'
        elif tab == 'SC':
            scope = r.get('AB', '').split('Scope: ', 1)[-1].split('. Command:', 1)[0]
            keys = ['boards']
            if scope.startswith('HD AFE'): keys = ['hd_afes']
            elif scope.startswith('AFE '): keys.append('afes')
            elif scope.startswith('Regulator'): keys.append('regulators')
            elif scope.startswith('Fan '): keys.append('fans')
            elif scope.startswith('ADC input'): keys.append('current_adc')
            expr, q = expression(keys), multiplier(keys)
            basis = scope
        elif tab == 'PoF':
            expr, q = expression(['pof_boxes']), multiplier(['pof_boxes'])
            basis = 'One instance of this specific temperature/voltage point per box; user 8 T + 8 V working case.'
        elif tab == 'SSP':
            keys = ['cal']
            if '{Channel}' in tag or tag == 'SSP.PulseMode': keys.append('cal_channels')
            elif '{Input}' in tag: keys.append('cal_voltages')
            expr, q = expression(keys), multiplier(keys)
            basis = r.get('AB', '').split('Scope: ', 1)[-1].split('. READ:', 1)[0]
        elif tab == 'WIENER':
            if tag == 'WIENER.PowerConsumption':
                expr, q = expression(['power_points']), multiplier(['power_points'])
                basis = 'Undefined measurement locations; kept as a separate addition.'
            else:
                expr, q = refs['boards'] + '+2*' + refs['cal'], values['boards'] + 2 * values['cal']
                basis = 'One DC output/board plus two outputs/calibration box; assigned outputs, not spare module capacity.'
        else:
            racks = values['top_racks'] + values['power_racks']
            rackexpr = '(' + expression(['top_racks', 'power_racks'], 'sum') + ')'
            if 'Outlet' in tag:
                expr = expression(['top_racks', 'top_outlets']) + '+' + expression(['power_racks', 'power_outlets'])
                q = multiplier(['top_racks', 'top_outlets']) + multiplier(['power_racks', 'power_outlets'])
                basis = '40 candidate 8-outlet top-plate PDUs + 4 standard 30-outlet PDUs; one PDU per rack is the draft allocation.'
            else:
                keys = ['temps'] if '.Temperature' in tag else ['leaks'] if '.Leak' in tag else ['smoke_hubs', 'contacts'] if '.Smoke' in tag else []
                expr = rackexpr + ('*' + expression(keys) if keys else '')
                q = racks * multiplier(keys)
                basis = 'Repeated reference rack profile; source proves reference hardware, not installed repetition.'
        add(tab, tag, role, expr, q, basis, status)

add('PoF', 'PoF.BoxEnable', 'Control', expression(['pof_boxes']), values['pof_boxes'],
    'User-requested single enable per box; additional proposal, absent from the implemented-variable inventory.', 'Proposed interface')
totals = defaultdict(lambda: dict(control=0, publish=0, total=0, unallocated_patterns=0))
for row in rows:
    t = totals[row['tab']]
    if row['quantity'] is None:
        t['unallocated_patterns'] += 1
    else:
        t[row['role'].lower()] += row['quantity']
        t['total'] += row['quantity']
grand = {key: sum(t[key] for t in totals.values()) for key in ['control', 'publish', 'total', 'unallocated_patterns']}
assert grand['total'] == 20804, grand
assert totals['SSP']['total'] == 137 * 8
assert totals['PDU']['total'] == 2*44 + 2*440 + 3*176 + 2*44 + 5*44
assert len(rows) == 160  # 159 unchanged catalogue definitions + proposed box enable.

wb = load_workbook(SOURCE)
inputs = wb.create_sheet('System Inputs')
inputs.append(['Input', 'Draft quantity', 'Meaning', 'Source / counting basis'])
for key, value, label, source in INPUTS:
    inputs.append([key, value, label, source])
validation = DataValidation(type='whole', operator='greaterThanOrEqual', formula1=0, allow_blank=True)
inputs.add_data_validation(validation); validation.add(f'B2:B{inputs.max_row}')
details = wb.create_sheet('System Expansion')
details.append(['Tab', 'Variable / pattern', 'Control / Publish', 'Counting basis', 'Read-path status', 'Expanded points', 'Allocated?'])
for i, row in enumerate(rows, 2):
    used = sorted(set(re.findall(r"'System Inputs'!\$B\$\d+", row['expression'])))
    expr = row['expression']
    formula = f'=IF(COUNT({",".join(used)})={len(used)},{expr},"")' if used else '=0'
    details.append([row['tab'], row['tag'], row['role'], row['basis'], row['status'], formula,
                    f'=IF(ISNUMBER(F{i}),"Counted","Quantity needed")'])
summary = wb.create_sheet('System Totals', 1)
summary.append(['FD-VD expanded points', 'Control', 'Publish', 'Subtotal', 'Unallocated patterns'])
for i, tab in enumerate(['DAQ', 'OPMON', 'SC', 'PoF', 'SSP', 'WIENER', 'PDU'], 2):
    last = details.max_row
    summary.append(['PDS calibration box' if tab == 'SSP' else tab,
        f'=SUMIFS(\'System Expansion\'!F$2:F${last},\'System Expansion\'!A$2:A${last},"{tab}",\'System Expansion\'!C$2:C${last},"Control")',
        f'=SUMIFS(\'System Expansion\'!F$2:F${last},\'System Expansion\'!A$2:A${last},"{tab}",\'System Expansion\'!C$2:C${last},"Publish")',
        f'=SUM(B{i}:C{i})',
        f'=COUNTIFS(\'System Expansion\'!A$2:A${last},"{tab}",\'System Expansion\'!G$2:G${last},"Quantity needed")'])
summary.append(['SUBTOTAL', '=SUM(B2:B8)', '=SUM(C2:C8)', '=SUM(D2:D8)', '=SUM(E2:E8)'])
summary.append(['Snapshot', grand['control'], grand['publish'], grand['total'], grand['unallocated_patterns']])
summary.append(['Additional points', '6 × total Hermes buffers + 12 × populated HD-mezzanine AFEs + defined WIENER power points'])
summary.append(['Complete total', '=IF(E9=0,D9,"Subtotal + unallocated points")'])
summary.append(['Case', 'FD-VD: 42 boards × 32 used channels, 4 AFE groups/board; two data links/board; 8 calibration boxes.'])
summary.append(['Rack profile', '40 top-plate + 4 PDS power racks; one PDU, four T, one leak and one smoke assembly per rack. Repetition is a draft assumption.'])
summary.append(['PDU inventory', '40 × 8 + 4 × 30 = 440 outlets; 176 T probes, 44 leak detectors, 44 smoke assemblies, 220 native contact-state fields.'])
summary.append(['Counting policy', 'Logical interface points: indexed fields expanded once; packed words count once; config and observed-status entries retain their distinct interface roles. Not a count of deployed OPC UA nodes.'])
summary.append(['Scope', 'PDS interfaces for FD-VD; excludes unrelated TPC, DAQ-server and central SC racks. Shared top-rack hardware is included once.'])
summary.append(['Implementation limits', 'V2 alternatives excluded from V3 case. Includes requested PoF and WIENER mappings and proposed box enable; subtotal is not an implemented/deployed claim.'])
summary.append(['Capacity alternative', '40 channels + 5 AFEs per board gives 23,828 before the same unallocated additions; the 32-channel case uses design active channels.'])
summary.append(['Catalogue', 'Original 159 definitions remain unchanged. The system estimate adds 80 proposed box-enable instances, separately from that catalogue.'])
summary.append(['PoF discrepancy', 'User working case is 8 T + 8 V; team presentation says 4–7 V + one chassis RTD. Do not present these as the same inventory.'])
main = wb['Main']; main['A1'] = 'PDS operational variables — DRAFT v1.2'
main['A25'] = 'FD-VD initial snapshot'; main['B25'] = '20,804 expanded points + unallocated additions; see System Totals for recalculation. The 159-definition table above is unchanged.'
main['A26'] = 'System Totals'; main['A26'].hyperlink = "#'System Totals'!A1"; main['A26'].style = 'Hyperlink'
for sheet in [inputs, details, summary]:
    sheet.freeze_panes = 'B2'; sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='203343'); cell.font = Font(color='FFFFFF', bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 42 if sheet != details else 60
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'; sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = '1:1'
for col, width in {'A':23, 'B':18, 'C':40, 'D':100}.items(): inputs.column_dimensions[col].width = width
for col, width in {'A':15, 'B':52, 'C':19, 'D':90, 'E':25, 'F':20, 'G':23}.items(): details.column_dimensions[col].width = width
for col, width in {'A':45, 'B':22, 'C':20, 'D':20, 'E':24}.items(): summary.column_dimensions[col].width = width
for row in range(11, 22):
    summary.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    summary.row_dimensions[row].height = 44
wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
wb.save(OUTPUT)

check = load_workbook(OUTPUT)
for name in ['DAQ', 'OPMON', 'SC', 'PoF', 'SSP', 'WIENER', 'PDU']:
    assert list(check[name].values) == list(load_workbook(SOURCE)[name].values), name
record = dict(scope='FD-VD explicit draft case; not final node inventory', source=str(SOURCE),
    source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(), workbook=str(OUTPUT),
    workbook_sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest(), inputs=INPUTS,
    tabs=dict(totals), totals=grand, additions='6 × B + 12 × M + W',
    rows=rows, original_variable_sheets_unchanged=True,
    pdu=dict(racks=44, pdus=44, outlets=440, temperatures=176, leaks=44, smoke_detectors=44, native_contact_fields=220),
    sources=dict(icd='/home/arroyave/work/icd-sc-pds/sections/02_physical_interface.tex',
        rack='/home/arroyave/work/racks/data/fd2-estimate-cells.csv',
        px3_guide='/home/arroyave/work/racks/sources/docdb-32655-v8/PX3_Rack_Protection_Setup.pdf',
        guide_locators='p.4 hardware; p.8 smoke; p.21 four Celsius sensors',
        pdu_udt='/home/arroyave/work/racks/sources/docdb-32655-v8/PDU_UDTs_2026_06_03.json',
        pdu_vendor_urls=[f'https://cdn.raritan.com/product-selector/pdus/{m}/{m}-spec.pdf' for m in ['PX3-5147R','PX3-5170V-E2F5']]))
(BASE/'system-counts.json').write_text(json.dumps(record, indent=2)+'\n')
tex = ['% Generated by build_system_counts.py.', f'\\newcommand{{\\SystemSubtotal}}{{{grand["total"]:,}}}',
       f'\\newcommand{{\\SystemControl}}{{{grand["control"]:,}}}', f'\\newcommand{{\\SystemPublish}}{{{grand["publish"]:,}}}',
       '\\newcommand{\\SystemCountsRows}{%']
for tab in ['DAQ','OPMON','SC','PoF','SSP','WIENER','PDU']:
    t = totals[tab]; label = 'PDS calibration box' if tab == 'SSP' else tab
    tex.append(f'{label} & {t["control"]:,} & {t["publish"]:,} & {t["total"]:,}' + r'\\')
tex.append('}'); (BASE/'system-counts.tex').write_text('\n'.join(tex)+'\n')
print(json.dumps(dict(workbook=str(OUTPUT), totals=grand, tabs=dict(totals))))
